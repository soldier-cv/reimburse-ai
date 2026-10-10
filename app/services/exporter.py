import io
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy.orm import Session
from ..models import ExpenseReport, Invoice, Project

def export_reports_to_excel(db: Session, project_id: int = 1) -> io.BytesIO:
    """
    根据数据库中的月度台账和发票数据，
    导出华总定制的【月度多行合并汇总版】Excel 文件流。
    """
    project = db.query(Project).filter(Project.id == project_id).first()
    proj_name = project.name if project else "上海饭范范科技有限公司"
    plate = project.plate_number if project else "苏U88M32"

    reports = db.query(ExpenseReport).filter(ExpenseReport.project_id == project_id).order_by(ExpenseReport.month.asc()).all()

    wb = openpyxl.Workbook()

    # 样式定义
    font_title = Font(name="微软雅黑", size=14, bold=True, color="1F4E79")
    font_sub = Font(name="微软雅黑", size=9.5, color="595959")
    font_th = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
    font_cell = Font(name="微软雅黑", size=10, color="000000")
    font_cell_bold = Font(name="微软雅黑", size=10, bold=True, color="000000")

    fill_th_main = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    fill_fuel_col = PatternFill(start_color="F2F7FA", end_color="F2F7FA", fill_type="solid")
    fill_fuel_calc = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid")
    fill_total_calc = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")

    thin_border_side = Side(style="thin", color="D9D9D9")
    med_border_bottom = Side(style="medium", color="1F4E79")
    border_cell = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

    # -----------------------------------------------------------------
    # Sheet 1: 费用汇总总览
    # -----------------------------------------------------------------
    ws1 = wb.active
    ws1.title = "费用汇总总览"
    ws1.views.sheetView[0].showGridLines = True

    ws1.merge_cells("A1:H1")
    ws1["A1"] = "交通差旅费用报销汇总表"
    ws1["A1"].font = font_title
    ws1["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws1.row_dimensions[1].height = 30

    ws1.merge_cells("A2:H2")
    ws1["A2"] = f"报销人：华旭东  |  车辆：{plate}  |  主体：{proj_name}  |  说明：通行费直接录入汇总单金额；燃油票默认预留3行逐笔填入（自动求和）"
    ws1["A2"].font = font_sub
    ws1["A2"].alignment = Alignment(horizontal="left", vertical="center")
    ws1.row_dimensions[2].height = 20

    headers = [
        "序号", "报销月份", "通行费 (元)\n[汇总单直填]", 
        "燃油票明细 (元)\n[每行填一张]", "燃油小计 (元)\n[公式求和]", 
        "月度总报销 (元)\n[公式合计]", "提交时间", "报销进度"
    ]
    ws1.row_dimensions[4].height = 28
    for c_idx, h in enumerate(headers, start=1):
        cell = ws1.cell(row=4, column=c_idx, value=h)
        cell.font = font_th
        cell.fill = fill_th_main
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border_cell

    current_row = 5
    seq = 1

    for rep in reports:
        # 查询该月的燃油发票
        fuel_invs = db.query(Invoice).filter(Invoice.month == rep.month, Invoice.category_code == "fuel").all()
        fuels = [inv.amount for inv in fuel_invs if inv.amount > 0]
        
        num_rows = max(3, len(fuels))
        start_row = current_row
        end_row = current_row + num_rows - 1

        for r in range(start_row, end_row + 1):
            ws1.row_dimensions[r].height = 22
            for c in range(1, 9):
                ws1.cell(row=r, column=c).border = border_cell

        # 序号
        ws1.merge_cells(start_row=start_row, start_column=1, end_row=end_row, end_column=1)
        c = ws1.cell(row=start_row, column=1, value=seq)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.font = font_cell

        # 报销月份
        ws1.merge_cells(start_row=start_row, start_column=2, end_row=end_row, end_column=2)
        c = ws1.cell(row=start_row, column=2, value=rep.month)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.font = font_cell_bold

        # 通行费
        ws1.merge_cells(start_row=start_row, start_column=3, end_row=end_row, end_column=3)
        c = ws1.cell(row=start_row, column=3, value=rep.toll_amount)
        c.number_format = "¥#,##0.00"
        c.alignment = Alignment(horizontal="right", vertical="center")
        c.font = font_cell

        # 燃油明细行
        for i in range(num_rows):
            r = start_row + i
            c_f = ws1.cell(row=r, column=4)
            c_f.fill = fill_fuel_col
            c_f.alignment = Alignment(horizontal="right", vertical="center")
            if i < len(fuels):
                c_f.value = fuels[i]
                c_f.number_format = "¥#,##0.00"
                c_f.font = font_cell
            else:
                c_f.number_format = "¥#,##0.00"

        # 燃油小计 (公式)
        ws1.merge_cells(start_row=start_row, start_column=5, end_row=end_row, end_column=5)
        c = ws1.cell(row=start_row, column=5, value=f"=SUM(D{start_row}:D{end_row})")
        c.number_format = "¥#,##0.00"
        c.alignment = Alignment(horizontal="right", vertical="center")
        c.font = font_cell_bold
        c.fill = fill_fuel_calc

        # 月度总报销 (公式)
        ws1.merge_cells(start_row=start_row, start_column=6, end_row=end_row, end_column=6)
        c = ws1.cell(row=start_row, column=6, value=f"=C{start_row}+E{start_row}")
        c.number_format = "¥#,##0.00"
        c.alignment = Alignment(horizontal="right", vertical="center")
        c.font = font_cell_bold
        c.fill = fill_total_calc

        # 提交时间
        ws1.merge_cells(start_row=start_row, start_column=7, end_row=end_row, end_column=7)
        c = ws1.cell(row=start_row, column=7, value=rep.submit_date or "")
        c.alignment = Alignment(horizontal="center", vertical="center")

        # 报销进度
        ws1.merge_cells(start_row=start_row, start_column=8, end_row=end_row, end_column=8)
        c = ws1.cell(row=start_row, column=8, value=rep.status or "待报销")
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.font = font_cell_bold

        for col in range(1, 9):
            cur = ws1.cell(row=end_row, column=col)
            cur.border = Border(
                left=cur.border.left,
                right=cur.border.right,
                top=cur.border.top,
                bottom=med_border_bottom
            )

        current_row = end_row + 1
        seq += 1

    # 数据验证
    dv = DataValidation(type="list", formula1='"待报销,待付款,已完成"', allow_blank=True)
    ws1.add_data_validation(dv)
    dv.add("H5:H100")

    widths1 = {'A': 8, 'B': 14, 'C': 18, 'D': 18, 'E': 17, 'F': 18, 'G': 15, 'H': 14}
    for col_l, w in widths1.items():
        ws1.column_dimensions[col_l].width = w
    ws1.freeze_panes = "C5"

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
